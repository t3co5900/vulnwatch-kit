"""Generación de Excel y PDF."""
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .analysis import LEVELS

COLORS = {"crítico": "C62828", "alto": "EF6C00", "medio": "F9A825", "bajo": "2E7D32"}
ORDER = {lvl: i for i, lvl in enumerate(reversed(LEVELS))}


def _sorted(vulns):
    return sorted(vulns, key=lambda v: (ORDER[v["level"]], -(v.get("score") or 0)))


def _affects(v):
    return {True: "SÍ", False: "NO", None: "Sin dato"}[v.get("affects_me")]


def write_excel(vulns, path, days):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Vulnerabilidades de los últimos {days} días — {datetime.now():%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Nivel", "Cantidad"])
    for lvl in reversed(LEVELS):
        ws.append([lvl.capitalize(), sum(1 for v in vulns if v["level"] == lvl)])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=COLORS[lvl])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color="FFFFFF")
    ws.column_dimensions["A"].width = 18

    ws2 = wb.create_sheet("Vulnerabilidades")
    head = ["CVE", "Nivel", "CVSS", "Productos", "Te afecta", "Explotada (CISA KEV)", "EPSS %",
            "Versión corregida", "Publicado", "Descripción", "Referencias", "Estado", "Notas"]
    ws2.append(head)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="37474F")
    for v in _sorted(vulns):
        ws2.append([
            v["id"], v["level"].capitalize(), v.get("score"), ", ".join(v["products"]), _affects(v),
            "SÍ" if v.get("kev") else "no", round((v.get("epss") or 0) * 100, 1),
            v.get("fixed", ""), v.get("published", "")[:10], v.get("description", ""),
            "\n".join(v.get("refs", [])), "Pendiente", "",
        ])
        cell = ws2.cell(ws2.max_row, 2)
        cell.fill = PatternFill("solid", fgColor=COLORS[v["level"]])
        cell.font = Font(bold=True, color="FFFFFF")
    for col, w in zip("ABCDEFGHIJKLM", [18, 11, 7, 22, 11, 14, 9, 18, 12, 80, 50, 12, 30]):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    wb.save(path)


def write_pdf(vulns, path, days):
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title="Reporte de vulnerabilidades")
    story = [Paragraph(f"Reporte de vulnerabilidades — últimos {days} días", styles["Title"]),
             Paragraph(f"Generado el {datetime.now():%d/%m/%Y %H:%M}. Total: {len(vulns)}", styles["Normal"]),
             Spacer(1, 0.4 * cm)]
    counts = [["Nivel", "Cantidad"]] + [[l.capitalize(), sum(1 for v in vulns if v["level"] == l)]
                                        for l in reversed(LEVELS)]
    t = Table(counts, colWidths=[4 * cm, 3 * cm])
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#37474F")),
             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
    for i, l in enumerate(reversed(LEVELS), start=1):
        style += [("BACKGROUND", (0, i), (0, i), colors.HexColor("#" + COLORS[l])),
                  ("TEXTCOLOR", (0, i), (0, i), colors.white)]
    t.setStyle(TableStyle(style))
    story += [t, Spacer(1, 0.5 * cm)]

    for lvl in reversed(LEVELS):
        group = [v for v in _sorted(vulns) if v["level"] == lvl]
        if not group:
            continue
        story.append(Paragraph(f"{lvl.capitalize()} ({len(group)})", styles["Heading2"]))
        rows = [["CVE", "CVSS", "Productos", "Te afecta", "KEV", "Corregida en", "Descripción"]]
        for v in group:
            desc = escape(v.get("description", ""))[:420]
            rows.append([Paragraph(escape(v["id"]), small), v.get("score") or "-",
                         Paragraph(escape(", ".join(v["products"])), small), _affects(v),
                         "SÍ" if v.get("kev") else "-", Paragraph(escape(v.get("fixed", "") or "-"), small),
                         Paragraph(desc, small)])
        tbl = Table(rows, colWidths=[3.2 * cm, 1.2 * cm, 3.2 * cm, 1.8 * cm, 1.2 * cm, 3 * cm, 12.5 * cm],
                    repeatRows=1)
        tbl.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + COLORS[lvl])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7.5)]))
        story += [tbl, Spacer(1, 0.4 * cm)]
    if not vulns:
        story.append(Paragraph("Sin novedades en el período.", styles["Normal"]))
    doc.build(story)
